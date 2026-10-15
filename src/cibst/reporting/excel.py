"""Excel results workbook for Finance and Credit stakeholders who work in spreadsheets."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cibst.reporting.insights import (migration_matrix, portfolio_snapshot, region_contribution, scenario_kpis,
                                      sector_contribution, top_obligors)

HEADER_FILL = PatternFill("solid", fgColor="1C5CAB")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(writer, df: pd.DataFrame, name: str, pct_cols=(), int_cols=()):
    df.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]
    for j, col in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=j)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        width = min(60, max(10, len(str(col)) + 2, *(len(f"{v}") for v in df[col].head(50).astype(str))))
        ws.column_dimensions[get_column_letter(j)].width = width
        fmt = "0.00%" if col in pct_cols else ("#,##0" if col in int_cols else None)
        if fmt is None and pd.api.types.is_float_dtype(df[col]):
            fmt = "#,##0.00"
        if fmt:
            for i in range(2, len(df) + 2):
                ws.cell(row=i, column=j).number_format = fmt
    ws.freeze_panes = "A2"


def write_workbook(arts, path: Path) -> None:
    snap = portfolio_snapshot(arts)
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        info = pd.DataFrame([["Run ID", arts.run_id], ["Reporting date", snap["as_of"]],
                             ["Data", "Synthetic, seeded (illustrative only)"],
                             ["Amounts", "USD millions"], ["Probability-weighted ECL t0", snap["weighted_ecl"]]],
                            columns=["Item", "Value"])
        _sheet(w, info, "README")
        k = scenario_kpis(arts)
        _sheet(w, k, "Scenario_KPIs", pct_cols=["cum_loss_rate", "stage2_peak", "rwa_uplift", "cet1_start", "cet1_trough"])
        paths = pd.concat([sc.paths.join(sc.regional_gdp.add_prefix("gdp_")).reset_index().assign(scenario=n)
                           for n, sc in arts.scenarios.scenarios.items()], ignore_index=True)
        _sheet(w, paths, "Scenario_paths")
        summ = pd.concat([r.summary.assign(scenario=n) for n, r in arts.results.items()], ignore_index=True)
        _sheet(w, summ, "Results_by_year", pct_cols=["coverage", "stage2_share", "stage3_share", "pd_wavg",
                                                     "pd12_fwd_wavg", "lgd_wavg", "rw_density", "cum_loss_rate", "cet1_ratio"])
        cap = pd.concat([r.capital.assign(scenario=n) for n, r in arts.results.items()], ignore_index=True)
        _sheet(w, cap, "Capital_walk", pct_cols=["cet1_ratio"])
        sev = "severe" if "severe" in arts.results else list(arts.results)[-1]
        _sheet(w, sector_contribution(arts, sev), "Sector_drivers",
               pct_cols=["impairment_rate", "share_of_impairment", "rwa_uplift", "stage2_share_peak"])
        _sheet(w, region_contribution(arts, sev), "Region_drivers", pct_cols=["impairment_rate"])
        _sheet(w, top_obligors(arts, sev, 25), "Top_obligors", pct_cols=["pd_t0", "pd_peak"])
        _sheet(w, migration_matrix(arts, sev).reset_index(), "Rating_migration",
               pct_cols=list(migration_matrix(arts, sev).columns))
        _sheet(w, arts.plausibility, "Plausibility_checks")
        if arts.sensitivity is not None:
            _sheet(w, arts.sensitivity, "Sensitivities")
        if arts.monitoring:
            _sheet(w, arts.monitoring["rag"], "Monitoring_RAG")
            _sheet(w, arts.monitoring["calibration"], "Calibration_by_grade", pct_cols=["avg_pd", "observed_dr"])
        _sheet(w, arts.dq.report, "DQ_report")
        _sheet(w, arts.controls, "Run_controls")
